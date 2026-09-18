#!/usr/bin/env python3
"""Build a research-only chromosome 21 portfolio workbook from existing VEP TSV.
Requires openpyxl; does not access the network or repeat variant calling/annotation.
"""
import csv
import pathlib
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path('chr21_analysis')
SOURCE = ROOT / 'results/annotation/ERR297183_chr21.vep_full.tsv'
DEST = ROOT / 'results/reports/ERR297183_chr21_portfolio.xlsx'
HEADERS = ['Variant_ID','Chromosome','Position_GRCh38','REF','ALT','QUAL','INFO_DP','DP4_REF','DP4_ALT','ALT_fraction_DP4','Gene','Gene_ID','Transcript','Transcript_selection','Consequence','Transcript_consequences','Impact_VEP','HGVSc','HGVSp','Review_reason','Review_status','IGV_notes','External_evidence','Limitations']
SEVERE = {'transcript_ablation','splice_acceptor_variant','splice_donor_variant','stop_gained','frameshift_variant','stop_lost','start_lost','transcript_amplification'}
MODERATE = {'missense_variant','inframe_insertion','inframe_deletion','protein_altering_variant','splice_donor_5th_base_variant','splice_region_variant'}
LIMITATIONS = 'chr21-only alignment may mis-map off-target reads; mean chr21 depth ~1.57x; research only; no independent validation'


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_info(value):
    result = {}
    for token in value.split(';'):
        if '=' in token:
            key, val = token.split('=', 1)
            result[key] = val
    return result


def is_true(value):
    return str(value).strip().lower() in ('1', 'yes', 'true')


def rank_transcript(row):
    return (bool(row.get('mane_select')), is_true(row.get('canonical')), bool(row.get('transcript_id')))


def excel_write(wb, name, headers, records):
    sheet = wb.create_sheet(name)
    sheet.append(headers)
    for record in records:
        sheet.append(record)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 19
    sheet.column_dimensions['A'].width = 33
    for index in range(20, min(len(headers), 24) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 35
    for index in (3, 6, 7, 8, 9, 10):
        if index <= len(headers):
            for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = '0.0000' if index == 10 else '0.##'
    return sheet


def main():
    if not SOURCE.is_file():
        raise SystemExit(f'Missing VEP TSV: {SOURCE}')
    grouped = defaultdict(list)
    with SOURCE.open(encoding='utf-8', newline='') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        required = {'input','gene_symbol','gene_id','transcript_id','most_severe_consequence','consequence_terms','impact','hgvsc','hgvsp'}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise SystemExit(f'Missing columns: {sorted(missing)}')
        for row in reader:
            fields = row['input'].split()
            if len(fields) < 8:
                raise SystemExit('Malformed VEP input; expected eight VCF columns')
            grouped[tuple(fields[:5])].append(row)
    if not grouped:
        raise SystemExit('No annotated variants found')

    all_rows, priority_rows = [], []
    for rows in grouped.values():
        best = max(rows, key=rank_transcript)
        chrom, pos, _, ref, alt, qual, _, info_text = best['input'].split()[:8]
        info = parse_info(info_text)
        dp4 = info.get('DP4', '').split(',')
        dp4_values = [number(v) for v in dp4] if len(dp4) == 4 else []
        valid_dp4 = len(dp4_values) == 4 and all(v is not None for v in dp4_values)
        ref_count = sum(dp4_values[:2]) if valid_dp4 else None
        alt_count = sum(dp4_values[2:]) if valid_dp4 else None
        fraction = round(alt_count / (ref_count + alt_count), 4) if valid_dp4 and ref_count + alt_count else None
        consequence = best.get('most_severe_consequence', '')
        terms = set(consequence.split('&')) | set(best.get('consequence_terms', '').replace('&', ',').split(','))
        quality = number(qual)
        dp = number(info.get('DP'))
        reason = ''
        if terms & SEVERE:
            reason = 'High-impact consequence; manual review required'
        elif terms & MODERATE:
            reason = 'Protein/splice-related consequence; manual review required'
        if reason and (quality is None or quality < 30 or dp is None or dp < 10 or alt_count is None or alt_count < 3):
            reason += '; read-support threshold not independently established'
        values = [f'{chrom}:{pos}:{ref}>{alt}', chrom, int(pos), ref, alt, quality, dp, ref_count, alt_count, fraction,
                  best.get('gene_symbol', ''), best.get('gene_id', ''), best.get('transcript_id', ''),
                  'MANE Select' if best.get('mane_select') else ('Canonical' if is_true(best.get('canonical')) else 'Other'),
                  consequence, best.get('consequence_terms', ''), best.get('impact', ''), best.get('hgvsc', ''), best.get('hgvsp', ''),
                  reason, 'Not reviewed', '', '', LIMITATIONS]
        all_rows.append(values)
        if reason:
            priority_rows.append(values[:])
    all_rows.sort(key=lambda row: (row[2], row[3], row[4]))
    priority_rows.sort(key=lambda row: (row[2], row[3], row[4]))

    wb = Workbook()
    wb.remove(wb.active)
    excel_write(wb, 'All_Variants', HEADERS, all_rows)
    excel_write(wb, 'Prioritized_Variants', HEADERS, priority_rows)
    review_headers = ['Variant_ID','Gene','Consequence','IGV_review','Read_support_notes','Mapping_concerns','ClinVar_gnomAD_notes','Review_status','Reviewer_notes']
    review_rows = [[row[0], row[10], row[14], 'Not reviewed', '', '', '', 'Pending', ''] for row in priority_rows]
    excel_write(wb, 'Variant_Review', review_headers, review_rows)
    notes = [
        ['Field', 'Description'],
        ['Project', 'ERR297183 chromosome 21 educational NGS portfolio'],
        ['Source', str(SOURCE)],
        ['Unique annotated variants', len(all_rows)],
        ['Selected for review', len(priority_rows)],
        ['Selection', 'Consequence-based triage only; not pathogenicity or clinical classification'],
        ['Limitations', 'Single-chromosome alignment can mis-map off-target reads; mean chr21 depth ~1.57x; no independent validation'],
        ['DP4 interpretation', 'ALT_fraction_DP4 = ALT/(REF+ALT) using high-quality DP4 counts; not an independently validated VAF'],
        ['Transcript display', 'MANE Select preferred, then canonical, then other; displayed transcript may not explain the most severe consequence across all transcripts'],
    ]
    excel_write(wb, 'README', notes[0], notes[1:])
    readme = wb['README']
    readme.column_dimensions['A'].width = 29
    readme.column_dimensions['B'].width = 105
    for row in readme.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
    DEST.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DEST)
    print(f'SUCCESS | unique variants: {len(all_rows)} | review candidates: {len(priority_rows)} | workbook: {DEST}')


if __name__ == '__main__':
    main()
